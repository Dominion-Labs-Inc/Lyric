#!/usr/bin/env python3
"""READ-01 — reading is the substrate's own sense, and one moment is one experience.

The substrate reads a document as it sees a picture and hears a recording: by its own sense, in the one act of
perceiving (`coord.perceive_moment`, through the `read` door). On the real substrate, in the SANDBOX:

  A  a text file is read by reading alone, its words go to the substrate's one reader, and what they state is
     kept with the line it was read from;
  B  the reading is remembered, with the runs of words it is known again by (never the words), and recorded in
     the reading ledger as a reading of the file as it is now;
  C  what the substrate's own document says is held as what it said -- an observation at the stated quality, not
     a truth asserted -- and the document `mentions` what it is about;
  D  a person's document is theirs: its memory is in their store, what it says goes to their context, and
     nothing of it enters the shared graph;
  E  the same text read again is known again (`same_text_as`); a new version of it is known as the same text
     changed; a different text is not known;
  F  a Word file and a file with no extension are read by what their bytes are;
  G  each door takes in what its sense can: `see` refuses a text and names `read`, `read` refuses a picture and
     names `see`;
  H  a text and a picture met at one moment are ONE memory, keeping what each sense keeps;
  J  a file a task's tool opens (`read_file`) is read by the substrate's own reading, as whoever the task is for:
     a person's task reads into their memory and their context;
  I  the main model's store is untouched.

Runs in the SANDBOX (`lyric_dev`) and does not empty it. Every memory it forms is deleted by id at the end. What
it leaves, named by this run's nonce: the nonce word the substrate's own document taught as a weak observation,
and what the test person's documents said, in that person's context -- the memory agent has no way to forget a
person's context.

    ./venv_lyric/bin/python3 experiments/READ-01/experiment.py
"""
from __future__ import annotations

import asyncio
import contextlib
import io
import os
import sys
import tempfile
import uuid
from pathlib import Path

os.environ.setdefault("POSTGRES_DATABASE", "lyric_dev")
for k, v in {"POSTGRES_PORT": "5433", "POSTGRES_USER": "stefan", "LYRIC_NO_WATCHDOG": "1",
             "TQDM_DISABLE": "1"}.items():
    os.environ.setdefault(k, v)
ROOT = Path(__file__).resolve().parents[1].parent
sys.path.insert(0, str(ROOT))

from experiments._evidence import RunRecord  # noqa: E402

NONCE = uuid.uuid4().hex[:6]
WORD = f"zorb{NONCE}"
PERSON = f"read01_{NONCE}"
PICTURE = str(ROOT / "test_data" / "vision_test.png")

EV = RunRecord(
    "READ-01",
    claim=("The substrate reads a document by its own sense, in the one act of perceiving: its words go to the "
           "substrate's one reader, what they state is held as what the document said (a person's in their "
           "context), the reading is remembered and recorded as a reading, the same text is known again by its "
           "runs of words, each door takes in only what its sense can, and a text and a picture met at one "
           "moment are one experience."),
    hypothesis=("A document read as a picture is seen (by sight's process, its words never read), a reading that "
                "forms no memory or no ledger entry, a person's document entering the shared graph, a text read "
                "twice not known again or a changed text taken for the same, a door that opens a thing its sense "
                "cannot take in, a moment remembered as two memories, or a write to the main store would each "
                "show here."))


def say(line=""):
    print(line, flush=True)


def check(name, ok, detail=""):
    EV.check(name, bool(ok), detail)
    say(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    return bool(ok)


async def main_store_rows() -> int:
    import asyncpg
    c = await asyncpg.connect(host="localhost", port=5433, user="stefan", database="lyric_db")
    try:
        n = 0
        for t in await c.fetch("select table_schema s, table_name n from information_schema.tables "
                               "where table_schema in ('unified','memory_hot','memory_cold') "
                               "and table_type='BASE TABLE'"):
            n += await c.fetchval(f'select count(*) from "{t["s"]}"."{t["n"]}"')
        return n
    finally:
        await c.close()


def stated(p):
    return [(f["subject"], f["relation"], f["object"]) for f in (p.content.get("stated") or [])] if p else []


async def main() -> int:
    main_before = await main_store_rows()
    quiet = io.StringIO()
    with contextlib.redirect_stdout(quiet), contextlib.redirect_stderr(quiet):
        from core.main import get_system
        system = get_system()
        await system.initialize()
        coord = system.autonomous_coordinator
        from core.database import get_database_manager
        db = get_database_manager()
        await db.assert_database_identity("lyric_dev")
    from core.agents.autonomous.autonomous_coordinator import Met
    work = Path(tempfile.mkdtemp(prefix="read01_"))
    formed = []

    def keep(p):
        mid = (p.metadata or {}).get("memory_id") if p is not None else None
        if mid:
            formed.append(mid)
        return mid

    def write(name, text):
        (work / name).write_text(text)
        return str(work / name)

    try:
        note = write("note.txt", f"A robin is a bird.\nA {WORD} is a bird.\nA snake is a reptile.\n")
        v2 = write("note_v2.txt", f"A robin is a bird.\nA {WORD} is a bird.\nA rose is a flower.\n")
        other = write("other.txt", "A cat is a mammal.\nA horse is a mammal.\nA lizard is a reptile.\n")
        mine = write("theirs.txt", f"A {WORD}y is a reptile.\nA cat is a mammal.\n")
        bare = write("notes", "A cat is a mammal.\n")
        import docx
        document = docx.Document()
        for line in ("A robin is a bird.", "A snake is a reptile."):
            document.add_paragraph(line)
        document.save(str(work / "note.docx"))
        worddoc = str(work / "note.docx")

        say("== A. A text read by reading, its words through the substrate's reader ==")
        own = await coord.read(note, actor_identity=None, source="read01")
        own_mid = keep(own)
        check("it is read by reading alone", own is not None and own.content.get("senses") == ["reading"],
              f"{own.content.get('senses') if own else None}")
        check("what its lines state is read by the substrate's reader",
              {("robin", "isa", "bird"), (WORD, "isa", "bird")} <= set(stated(own)), f"{stated(own)}")
        check("each fact is kept with the line it was read from",
              all(f.get("said") for f in own.content.get("stated") or []) if own else False)

        say("\n== B. The reading is remembered and recorded ==")
        media = await coord.recall_media(own_mid) if own_mid else []
        check("the reading forms a memory, keeping the runs of words it is known again by",
              bool(own_mid) and [m["perceived"].get("kind") for m in media] == ["text_trace"],
              f"{[m['perceived'].get('kind') for m in media]}")
        check("the file itself is not kept", all(not str(m["mime"]).startswith("text/") for m in media))
        check("the reading ledger holds it as read now", coord.reading.current(note) is not None)

        say("\n== C. What its own document says is held as what it said ==")
        rows = await db.execute_query(
            "SELECT source_concept_id, relation, target_concept_id FROM unified.concept_relations "
            "WHERE source_concept_id ILIKE $1 OR target_concept_id ILIKE $1", (f"%{WORD}%",), fetch_all=True) or []
        pairs = {(str(r["source_concept_id"]).split(":")[-1], str(r["relation"]),
                  str(r["target_concept_id"]).split(":")[-1]) for r in rows}
        check("what it states is in the shared graph", any(s == WORD and r == "isa" for s, r, _ in pairs),
              f"{sorted(pairs)}")
        check("the document mentions what it is about", any(r == "mentions" and o == WORD for _, r, o in pairs),
              f"{sorted(pairs)}")
        belief = coord.learning.belief_for_claim(f"{WORD} isa bird")
        EV.metric("belief in what the document said",
                  getattr(belief, "posterior", getattr(belief, "probability", None)),
                  note="held at the stated quality, an observation")

        say("\n== D. A person's document is theirs ==")
        theirs = await coord.read(mine, actor_identity=PERSON, source="read01")
        their_mid = keep(theirs)
        item = await coord.memory.postgres_storage.get_memory(their_mid) if their_mid else None
        check("its memory is theirs", item is not None and (item.user_id or None) == PERSON,
              f"{getattr(item, 'user_id', None)}")
        shared = await db.execute_query(
            "SELECT count(*) AS n FROM unified.concept_relations WHERE source_concept_id ILIKE $1",
            (f"%{WORD}y%",), fetch_one=True)
        scoped = await db.execute_query(
            "SELECT count(*) AS n FROM unified.scoped_concept_relations WHERE scope_actor = $1 AND subj ILIKE $2",
            (PERSON, f"%{WORD}y%"), fetch_one=True)
        check("what it says is held in their context", (scoped or {}).get("n", 0) >= 1, f"{scoped}")
        check("nothing of it enters the shared graph", (shared or {}).get("n", 1) == 0, f"{shared}")

        say("\n== E. The same text is known again; a changed one as changed; another not at all ==")
        again = await coord.read(note, actor_identity=None, source="read01")
        keep(again)
        before = again.content.get("read_before") or [] if again else []
        check("read again, it is known as the same text", any(e.get("same_text") for e in before),
              f"{[(e.get('subject'), e.get('support'), e.get('same_text')) for e in before]}")
        changed = await coord.read(v2, actor_identity=None, source="read01")
        keep(changed)
        seen = changed.content.get("read_before") or [] if changed else []
        check("a new version is known as that text, changed",
              bool(seen) and not any(e.get("same_text") for e in seen),
              f"{[(e.get('subject'), e.get('support'), e.get('same_text')) for e in seen]}")
        EV.metric("share a new version agreed with", max((e.get("support") or 0) for e in seen) if seen else 0.0)
        fresh = await coord.read(other, actor_identity=None, source="read01")
        keep(fresh)
        check("a different text is not known", fresh is not None and not fresh.content.get("read_before"),
              f"{fresh.content.get('read_before') if fresh else None}")

        say("\n== F. Read by what the bytes are ==")
        word = await coord.read(worddoc, actor_identity=PERSON, source="read01")
        keep(word)
        check("a Word file is read, paragraph by paragraph",
              word is not None and word.content.get("pages") == 2 and ("robin", "isa", "bird") in stated(word),
              f"{word.content.get('caption') if word else None}; {stated(word)}")
        named_by_bytes = await coord.read(bare, actor_identity=PERSON, source="read01")
        keep(named_by_bytes)
        check("a file with no extension is read as the text it is",
              ("cat", "isa", "mammal") in stated(named_by_bytes), f"{stated(named_by_bytes)}")

        say("\n== G. Each door takes in what its sense can ==")
        refusals = {}
        for door, path in (("see", note), ("read", PICTURE)):
            try:
                keep(await getattr(coord, door)(path, actor_identity=PERSON, source="read01"))
                refusals[door] = ""
            except ValueError as error:
                refusals[door] = str(error)
        check("`see` refuses a text and names `read`", refusals["see"].endswith("use read"), refusals["see"])
        check("`read` refuses a picture and names `see`", refusals["read"].endswith("use see"), refusals["read"])

        say("\n== H. A text and a picture at one moment: one experience ==")
        both = await coord.perceive_moment([Met(note, "read"), Met(PICTURE, "see")],
                                           actor_identity=PERSON, source="read01")
        ids = [keep(p) for p in both]
        kinds = sorted(m["perceived"].get("kind") for m in await coord.recall_media(ids[0])) if ids[0] else []
        check("both are taken in, each by its own sense",
              [p.content.get("senses") if p else None for p in both] == [["reading"], ["sight"]])
        check("they are one memory", len(set(ids)) == 1 and ids[0] is not None, f"{ids}")
        check("it keeps what each sense keeps", kinds == ["sight_trace", "text_trace"], f"{kinds}")

        say("\n== J. What a task opens, the substrate reads ==")
        from core.agents.autonomous.shared_types import Priority, Task, TaskSource, TaskType
        task_note = write("task_note.txt", f"A {WORD}t is a mammal.\n")
        task = Task(id=f"read01_task_{NONCE}", type=TaskType.ANALYSIS, description="Read my note.",
                    priority=Priority.HIGH, source=TaskSource.MANUAL, actor=PERSON, metadata={})
        aware_before = len(coord.vision.recent_percepts(1000))
        ran = await coord._run_tool("read_file", {"file_path": task_note}, task)
        check("the task's tool opened the file", bool(ran) and ran.get("success"), f"{(ran or {}).get('error')}")
        read_now = [p for p in coord.vision.recent_percepts(1000)[aware_before:]
                    if (p.content or {}).get("senses") == ["reading"]]
        task_mid = keep(read_now[-1]) if read_now else None
        task_item = await coord.memory.postgres_storage.get_memory(task_mid) if task_mid else None
        check("the substrate read it, by its own reading", bool(read_now)
              and (f"{WORD}t", "isa", "mammal") in stated(read_now[-1]), f"{stated(read_now[-1]) if read_now else None}")
        check("the reading is the person's memory", task_item is not None and (task_item.user_id or None) == PERSON,
              f"{getattr(task_item, 'user_id', None)}")
        task_scoped = await db.execute_query(
            "SELECT count(*) AS n FROM unified.scoped_concept_relations WHERE scope_actor = $1 AND subj = $2",
            (PERSON, f"{WORD}t"), fetch_one=True)
        check("what it says is in their context", (task_scoped or {}).get("n", 0) >= 1, f"{task_scoped}")
        check("the reading ledger holds it", coord.reading.current(task_note) is not None)
    finally:
        gone = 0
        for mid in dict.fromkeys(formed):
            gone += bool(await coord.memory.delete_memory(mid, reason="READ-01"))
        say(f"\n    memories formed and deleted: {gone}")
        with contextlib.redirect_stdout(quiet), contextlib.redirect_stderr(quiet):
            coord.vision.close()
        for f in work.iterdir():
            f.unlink()
        work.rmdir()

    say("\n== I. The main model ==")
    main_after = await main_store_rows()
    check("the main model's store is untouched", main_before == main_after, f"{main_before} rows before, {main_after} after")
    EV.note(f"nonce {NONCE}: left in the sandbox, the word {WORD} (a weak observation from the substrate's own "
            f"document) and person {PERSON}'s context")
    await EV.verify_database()
    EV.write()
    passed = sum(1 for c in EV.checks if c.passed)
    say(f"\n{passed}/{len(EV.checks)} passed")
    return 0 if passed == len(EV.checks) else 1


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
