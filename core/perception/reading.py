#!/usr/bin/env python3
"""Reading: the substrate's sense for written words, as sight is for pixels and
hearing for sound.

What runs here is only what reading MEASURES of a file: its pages of text, what
the file really is (its bytes say, not its name), and the trace a memory keeps
of it so the same text is known when it is met again. What the words MEAN is
not measured here: they go to the substrate's one reader
(`core.semantics.derived_reader`), the same reader every sentence it is told
goes to, spoken or typed.

A TEXT IS KNOWN AGAIN BY ITS RUNS OF WORDS, as a sound is by its landmarks and
a picture by its keypoints. Every run of `SHINGLE_WORDS` consecutive words is
hashed, so two copies of a text share every run wherever they start, and a text
read inside a longer one (a quoted page, a new version of a document) shares the
runs of the part they have in common. The words themselves are not kept: what a
memory keeps of a reading is what it was known by and what it said.
"""
from __future__ import annotations

import hashlib
import io
import re
import zlib
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

#: WHAT A DOCUMENT IS, the file says, not its name: its first bytes where they
#: declare a format. Eight files in one real folder were misnamed; one was 292
#: bytes. Dispatching on the extension trusted the filename over the file.
_MAGIC: List[tuple] = [
    (b"%PDF", "pdf"),
    (b"PK\x03\x04", "zip"),          # docx/xlsx are zip containers
    (b"{\\rtf", "text"),
    (b"<!DOC", "text"), (b"<!doc", "text"),
    (b"<html", "text"), (b"<HTML", "text"), (b"<?xml", "text"),
]

DOCUMENT_READERS: Dict[str, str] = {
    ".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx",
    ".txt": "text", ".md": "text", ".csv": "text", ".json": "text",
    ".log": "text", ".yaml": "text", ".yml": "text",
}

#: How much of a file's start is looked at to tell text from binary.
_SNIFF_BYTES = 8192


def declared_kind(raw: bytes) -> Optional[str]:
    """What the FILE says it is, or None where its bytes do not say.

    Plain text declares nothing -- a .txt, .md, .csv or .log has no magic -- so
    None means "the bytes are silent", and the extension is then the only thing
    anyone knows. That is not a fallback: it is the honest order of evidence,
    content first and the name only where content does not speak.
    """
    # LEADING BLANK LINES DO NOT UNDECLARE A FILE. One of the eight was HTML
    # behind four newlines, so an 8-byte window missed `<!DOCTYPE` and the file
    # went to the PDF reader anyway. A binary format has no leading whitespace,
    # so stripping it costs those nothing and rescues the text ones.
    head = raw[:64].lstrip()[:8]
    for magic, kind in _MAGIC:
        if head.startswith(magic):
            return kind
    return None


def is_text(raw: bytes) -> bool:
    """Whether these bytes are written words: no NUL in the first stretch, and
    that stretch decodes as UTF-8 (a character cut at the end of it is not a
    reason to call a text binary). What the bytes are, the bytes say."""
    head = raw[:_SNIFF_BYTES]
    if b"\x00" in head:
        return False
    try:
        head.decode("utf-8")
        return True
    except UnicodeDecodeError as error:
        return error.start >= len(head) - 3


def kind_of(path: str) -> Optional[str]:
    """What reading would open a file as -- "pdf", "docx", "xlsx" or "text" --
    or None when it holds no written words reading can open. Its bytes decide
    first; its name only where they are silent."""
    p = Path(path)
    named = DOCUMENT_READERS.get(p.suffix.lower())
    try:
        with open(p, "rb") as f:
            raw = f.read(_SNIFF_BYTES)
    except OSError:
        return named
    declared = declared_kind(raw)
    if declared == "zip":
        return named if named in ("docx", "xlsx") else None
    if declared is not None:
        return declared
    if named is not None:
        return named
    return "text" if raw and is_text(raw) else None


def read_document(path: str) -> Dict[str, Any]:
    """A document read into its pages, with what it really is: `kind`, `named`
    (what its name said, when that differs), `sha256`, `pages`, and `lossy`
    when its text was not valid UTF-8. A file with no written words reading can
    open raises."""
    p = Path(path)
    raw = p.read_bytes()
    named = DOCUMENT_READERS.get(p.suffix.lower())
    kind = named
    declared = declared_kind(raw)
    if declared == "zip":
        # A zip container is a .docx or an .xlsx; which one only the extension
        # distinguishes, and here the two agree often enough that the name is
        # the evidence available.
        declared = kind if kind in ("docx", "xlsx") else None
    if declared:
        kind = declared
    if kind is None and raw and is_text(raw):
        kind = "text"
    if kind is None:
        raise ValueError(
            f"no reader for {p.name}: it holds no written words the substrate can "
            f"open, and it will not guess at its contents")
    lossy = False
    if kind == "pdf":
        from pypdf import PdfReader
        pages = [(page.extract_text() or "").strip() for page in PdfReader(io.BytesIO(raw)).pages]
    elif kind == "docx":
        import docx
        # A .docx has no pages until it is laid out; its paragraphs are the
        # structure it really has, so they are what is reported.
        pages = [para.text.strip() for para in docx.Document(io.BytesIO(raw)).paragraphs
                 if para.text.strip()]
    elif kind == "xlsx":
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        pages = []
        for sheet in book.worksheets:
            rows = [" ".join(str(c) for c in row if c is not None)
                    for row in sheet.iter_rows(values_only=True)]
            body = "\n".join(r for r in rows if r.strip())
            if body:
                pages.append(f"{sheet.title}\n{body}")
        book.close()
    else:
        # Encoding is DECLARED, never guessed silently: an undecodable byte is
        # replaced and the fact is reported, because a mangled character in a
        # mission document is a wrong reading and should be visible.
        try:
            pages = [raw.decode("utf-8")]
        except UnicodeDecodeError:
            pages, lossy = [raw.decode("utf-8", errors="replace")], True
    return {"kind": kind, "named": named if named and named != kind else None,
            "sha256": hashlib.sha256(raw).hexdigest(), "pages": pages, "lossy": lossy}


# ── the trace a reading keeps ───────────────────────────────────────────────

#: A run of this many words is what a text is known again by.
SHINGLE_WORDS = 5
#: The most runs a trace keeps: the smallest hashes, so any two traces of one
#: text keep the same ones. A text of about this many words keeps every run.
TEXT_MAX_KEYS = 8192
#: READ BEFORE. A text met again shares at least this many kept runs with the
#: reading remembered, and at least this share of the shorter of the two.
READ_MIN_SHARED = 3
READ_MIN_SHARE = 0.3

_WORD = re.compile(r"[^\W_]+(?:['’][^\W_]+)*", re.UNICODE)


def words_of(pages: Iterable[str]) -> List[str]:
    """The words of a text, in order, lowercased: what its runs are made of."""
    return [w.lower() for page in pages for w in _WORD.findall(page or "")]


def shingles(pages: Iterable[str]) -> List[int]:
    """The runs a text is known again by: each run of `SHINGLE_WORDS` words,
    hashed to a signed 32-bit key (the store keeps keys as integers), smallest
    first, at most `TEXT_MAX_KEYS`. A text shorter than one run keeps its words
    as one run."""
    words = words_of(pages)
    if not words:
        return []
    runs = ([" ".join(words)] if len(words) < SHINGLE_WORDS else
            (" ".join(words[i:i + SHINGLE_WORDS]) for i in range(len(words) - SHINGLE_WORDS + 1)))
    kept = {zlib.crc32(run.encode("utf-8")) - (1 << 31) for run in runs}
    return sorted(kept)[:TEXT_MAX_KEYS]


def trace_bytes(keys: Sequence[int]) -> Optional[bytes]:
    """What a memory keeps of a reading: its kept runs, as a numpy archive (as a
    hearing's trace is). None for a text with no words."""
    if not keys:
        return None
    import numpy as np
    buffer = io.BytesIO()
    np.savez_compressed(buffer, shingles=np.asarray(keys, np.int32))
    return buffer.getvalue()


def trace_keys(data: bytes):
    """A kept reading's runs, back from its trace; None when it is not one."""
    import numpy as np
    try:
        with np.load(io.BytesIO(bytes(data))) as archive:
            if "shingles" not in archive.files:
                return None
            return np.asarray(archive["shingles"], np.int32)
    except Exception:
        return None


def agreement(met: Sequence[int], kept: Sequence[int]) -> Tuple[int, float, bool]:
    """How much a text met now shares with one read before: the kept runs they
    share, what share of the shorter of the two that is, and whether they are
    the same text -- every run of each in the other. A page of a document shares
    all of its own runs with the document and is still not the document."""
    a, b = set(int(k) for k in met), set(int(k) for k in kept)
    if not a or not b:
        return 0, 0.0, False
    shared = len(a & b)
    return shared, shared / min(len(a), len(b)), shared == len(a) == len(b)
